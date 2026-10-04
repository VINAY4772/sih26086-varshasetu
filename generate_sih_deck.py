import os
import sys
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Path constants
SRC_PPTX = '/Users/maradanasaikiran/Downloads/Smart India Hackathon 2026 - VarshaSetu Presentation (1).pptx'
OUT_DIR = '/Users/maradanasaikiran/vinay new sih/SIH_PPT'
OUT_PPTX = os.path.join(OUT_DIR, 'VarshaSetu_SIH2026_6Slide_Official.pptx')
IMG_DIR = '/Users/maradanasaikiran/vinay new sih/varshshetu full doc/img'

os.makedirs(OUT_DIR, exist_ok=True)
prs = Presentation(SRC_PPTX)

print(f"Loaded presentation with {len(prs.slides)} slides.")
